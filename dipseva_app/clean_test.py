import openpyxl
p = "/home/shaarx/Documents/Dipseva/Dipseva_Ledger_NEW.xlsx"
wb = openpyxl.load_workbook(p)
ws = wb["Ledger"]
for r in range(3, 1003):
    if ws.cell(r, 2).value == "TEST CHECK":
        print("clearing row", r)
        for c in range(1, 12):
            ws.cell(r, c).value = None
wb.save(p)
print("cleaned")
