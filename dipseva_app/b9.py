from b8 import *
from datetime import datetime
from openpyxl.utils.datetime import from_excel
lr=3; fixed=[]
for r in range(2,sS.max_row+1):
    dt=sS.cell(r,1).value
    if dt is None or (isinstance(dt,str) and dt.strip()==""): continue
    if isinstance(dt,datetime): d=dt.date()
    elif isinstance(dt,(int,float)):
        try: d=from_excel(dt).date()
        except Exception: continue
    else: continue
    item=sS.cell(r,2).value
    item=str(item).strip() if item is not None and str(item).strip()!="" else ""
    cust=sS.cell(r,7).value
    cust=str(cust).strip() if cust is not None and str(cust).strip()!="" else "CASH"
    def num(x):
        try: return float(x or 0)
        except Exception: return 0.0
    qty=num(sS.cell(r,4).value); rate=num(sS.cell(r,5).value)
    paid=num(sS.cell(r,9).value); ret=num(sS.cell(r,11).value)
    narr=str(sS.cell(r,13).value or "").strip()
    if item=="":
        if paid:
            ws.cell(row=lr,column=1,value=d).number_format=DF
            ws.cell(row=lr,column=2,value=cust)
            ws.cell(row=lr,column=3,value="Payment")
            ws.cell(row=lr,column=9,value=paid).number_format=MF
            ws.cell(row=lr,column=11,value=("Advance "+narr).strip())
            lr+=1
        continue
    ws.cell(row=lr,column=1,value=d).number_format=DF
    ws.cell(row=lr,column=2,value=cust)
    ws.cell(row=lr,column=3,value="Issue")
    ws.cell(row=lr,column=4,value=item)
    ws.cell(row=lr,column=6,value=qty).number_format='#,##0'
    ws.cell(row=lr,column=7,value=rate).number_format=MF
    f_old=sS.cell(r,6).value
    try: fo=float(f_old) if f_old not in (None,"") and not (isinstance(f_old,str) and f_old.startswith("=")) else qty*rate
    except Exception: fo=qty*rate
    if abs(fo-qty*rate)>0.01:
        fixed.append((r,item,cust,qty,rate,fo,qty*rate))
        narr=(f"OLD {fo:g} FIX->{qty*rate:g} "+narr).strip()
    ws.cell(row=lr,column=9,value=paid if paid else None)
    if paid: ws.cell(row=lr,column=10,value="Cash")
    ws.cell(row=lr,column=11,value=narr)
    lr+=1
    if ret>0 and ret<qty:
        ws.cell(row=lr,column=1,value=d).number_format=DF
        ws.cell(row=lr,column=2,value=cust)
        ws.cell(row=lr,column=3,value="Return")
        ws.cell(row=lr,column=4,value=item)
        ws.cell(row=lr,column=6,value=ret).number_format='#,##0'
        ws.cell(row=lr,column=11,value="old Return")
        lr+=1
    if lr>N: break
print("ledger",lr-3); print("fixed",fixed)
wb.active=wb["Dashboard"]
wb.save(OUT)
print("saved",OUT)
