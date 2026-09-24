import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
NAVY="1B2A4E"; SAF="D97706"; TEAL="0D9488"; GRN="16A34A"; RED="DC2626"
HF=Font(name="Noto Sans Devanagari",size=11,bold=True,color="FFFFFF")
HFL=PatternFill("solid",fgColor=NAVY)
BF=Font(name="Noto Sans Devanagari",size=11,color="0F172A")
MF='#,##0'; DF='DD-MMM-YYYY'
TH=Side(style="thin",color="E2E8F0"); BD=Border(left=TH,right=TH,top=TH,bottom=TH)
CE=Alignment(horizontal="center",vertical="center",wrap_text=True)
LF=Alignment(horizontal="left",vertical="center",wrap_text=True)
N=1002
OUT="/home/shaarx/Documents/Dipseva/Dipseva_Ledger_NEW.xlsx"
SRC="/home/shaarx/Documents/Dipseva/Ledger JULY 2026.xlsx"
wb=openpyxl.Workbook()
ws=wb.active; ws.title="Ledger"
ws.sheet_properties.pageSetUpPr.fitToPage=True
ws.freeze_panes="A2"; ws.sheet_properties.tabColor=NAVY
