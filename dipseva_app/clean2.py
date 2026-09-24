import openpyxl
p = "/home/shaarx/Documents/Dipseva/Dipseva_Ledger_NEW.xlsx"
wb = openpyxl.load_workbook(p)
ws = wb["Ledger"]
n = 0
for r in range(3, 1003):
    if ws.cell(r, 2).value == "ZZ TMP":
        for c in range(1, 12):
            ws.cell(r, c).value = None
        ws.cell(r, 5).value = f'=IF($D{r}="","",IFERROR(XLOOKUP($D{r},Items!$A:$A,Items!$B:$B),"ERR"))'
        ws.cell(r, 7).value = f'=IF($D{r}="","",IFERROR(XLOOKUP($D{r},Items!$A:$A,Items!$D:$D),0))'
        ws.cell(r, 8).value = f'=IF($C{r}="Issue",IFERROR($F{r}*$G{r},0),0)'
        n += 1
wb.save(p)
print("removed ZZ TMP rows:", n)
