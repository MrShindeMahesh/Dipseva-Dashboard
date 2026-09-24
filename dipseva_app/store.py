"""Workbook IO + summaries for Dipseva."""
import os
import re
from datetime import datetime, date, timedelta
import openpyxl
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(BASE, "Dipseva_Ledger_NEW.xlsx")
LEDGER_FIRST, LEDGER_LAST = 3, 1002
DATE_FMT, MONEY_FMT = "DD-MMM-YYYY", "#,##0"
MODES = ["Cash", "UPI", "PhonePe SK", "Paytm SK", "PhonePe VK", "PAYTM B QR", "Bank"]
TYPES = ["Issue", "Return", "Payment"]


def load_wb():
    return openpyxl.load_workbook(XLSX)


def money(v, default=0.0):
    """Tolerant money/qty parser — accepts ₹1150, '1,150', '1 150.50', '-', ''. Never raises."""
    if v is None or isinstance(v, bool):
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("₹", "")
    s = s.replace("Rs.", "").replace("Rs", "").replace("INR", "").replace(" ", "")
    s = s.replace(",", "").replace("_", "")
    if s in ("", "-", "--", "."):
        return default
    try:
        return float(s)
    except Exception:
        hit = re.findall(r"-?\d+(?:\.\d+)?", s)
        return float(hit[0]) if hit else default


def parse_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except Exception:
            pass
    try:
        return datetime.fromisoformat(str(v)).date()
    except Exception:
        return None
def row_excel(r):
    """Excel formulas every Ledger row keeps (Item name lookup, rate lookup, bill)."""
    return {
        5: f'=IF($D{r}="","",IFERROR(XLOOKUP($D{r},Items!$A:$A,Items!$B:$B),"ERR"))',
        7: f'=IF($D{r}="","",IFERROR(XLOOKUP($D{r},Items!$A:$A,Items!$D:$D),0))',
        8: f'=IF($C{r}="Issue",IFERROR($F{r}*$G{r},0),0)',
    }


def next_free_row(ws):
    """First empty Ledger row (A+B+C all blank), or None when the sheet is full."""
    r = LEDGER_FIRST
    while r <= LEDGER_LAST:
        if all(ws.cell(r, c).value in (None, "") for c in (1, 2, 3)):
            return r
        r += 1
    return None


def write_line(ws, r, dt, cust, typ, item="", qty=0, rate=0, paid=0, mode="", note=""):
    """Write one ledger line. Item lines keep qty/rate + Excel formulas; Payment keeps paid."""
    for c in range(1, 12):
        ws.cell(r, c).value = None
    ws.cell(r, 1).value = dt
    ws.cell(r, 1).number_format = DATE_FMT
    ws.cell(r, 2).value = cust
    ws.cell(r, 3).value = typ
    if typ in ("Issue", "Return"):
        f = row_excel(r)
        ws.cell(r, 4).value = item
        ws.cell(r, 5).value = f[5]
        ws.cell(r, 6).value = qty or None
        ws.cell(r, 6).number_format = "#,##0"
        ws.cell(r, 7).value = rate if rate else f[7]
        ws.cell(r, 7).number_format = MONEY_FMT
        ws.cell(r, 8).value = f[8]
        ws.cell(r, 8).number_format = MONEY_FMT
    ws.cell(r, 9).value = paid or None
    ws.cell(r, 9).number_format = MONEY_FMT
    ws.cell(r, 10).value = (mode or "Cash") if paid else (mode or None)
    ws.cell(r, 11).value = note or None
    return r


def clear_row(ws, r):
    """Blank a ledger row but keep the Excel formulas so the sheet still works by hand."""
    for c in range(1, 12):
        ws.cell(r, c).value = None
    for c, f in row_excel(r).items():
        ws.cell(r, c).value = f


def read_ledger():
    wb = load_wb()
    ws = wb["Ledger"]
    wi = wb["Items"]
    imap = {}
    for rr in range(3, 203):
        k = wi.cell(rr, 1).value
        if k not in (None, ""):
            imap[str(k).strip()] = str(wi.cell(rr, 2).value or "").strip()
    rows = []
    for r in range(3, 1003):
        a = ws.cell(r, 1).value; b = ws.cell(r, 2).value; c = ws.cell(r, 3).value
        if (a in (None, "")) and (b in (None, "")) and (c in (None, "")):
            continue
        d = ws.cell(r, 4).value; f = ws.cell(r, 6).value; g = ws.cell(r, 7).value
        i = ws.cell(r, 9).value
        try: qty = float(f or 0)
        except Exception: qty = 0
        try: rate = float(g or 0)
        except Exception: rate = 0
        typ = str(c or "").strip()
        bill = qty * rate if typ == "Issue" else 0.0
        try: paid = float(i or 0)
        except Exception: paid = 0.0
        dt = parse_date(a)
        item = str(d or "").strip()
        if item.startswith("="):
            item = ""
        ev = ws.cell(r, 5).value
        es = str(ev or "").strip()
        if es.startswith("=") or es in ("", "ERR"):
            iname = imap.get(item, "")
        else:
            iname = es
        rows.append({"row": r, "date": dt.isoformat() if dt else "",
            "date_disp": dt.strftime("%d-%b") if dt else "",
            "customer": str(b or "").strip(), "type": typ,
            "item": item,
            "item_name": iname,
            "qty": int(qty) if qty == int(qty) else qty,
            "rate": rate, "bill": bill, "paid": paid,
            "mode": str(ws.cell(r, 10).value or "").strip(),
            "note": str(ws.cell(r, 11).value or "").strip()})
    return rows
def read_masters():
    wb = load_wb()
    wi = wb["Items"]; wc = wb["Customers"]
    items = []
    for r in range(3, 203):
        a = wi.cell(r, 1).value
        if a in (None, ""):
            continue
        items.append({"id": str(a).strip(), "name": str(wi.cell(r, 2).value or "").strip(),
            "marathi": str(wi.cell(r, 3).value or "").strip(),
            "rate": wi.cell(r, 4).value or 0, "opening": wi.cell(r, 5).value or 0,
            "min": wi.cell(r, 6).value or 0})
    custs = []
    for r in range(3, 203):
        a = wc.cell(r, 1).value
        if a in (None, ""):
            continue
        custs.append({"name": str(a).strip(), "person": str(wc.cell(r, 2).value or "").strip(),
            "phone": str(wc.cell(r, 3).value or "").strip(),
            "ctype": str(wc.cell(r, 4).value or "").strip(),
            "addr": str(wc.cell(r, 5).value or "").strip(),
            "opening_due": wc.cell(r, 6).value or 0})
    return items, custs
