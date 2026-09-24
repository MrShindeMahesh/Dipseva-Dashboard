from b7 import *
import openpyxl as _ox
src=_ox.load_workbook(SRC,data_only=True)
sS=src["Sales"]; sI=src["Inventory"]; sC=src["Customers"]
MR={"CTRNG":"केटरिंग","Chair ":"खुर्ची","Round Table ":"गोल टेबल","Fan":"पंखा","Gadi ":"गादी","Drum":"ड्रम","Table":"टेबल","Ahuja Speaker":"अहुजा स्पीकर","Samai":"समई","Wash Basin":"वॉश बेसिन","Bedsheet ":"चादर","Palna":"पाळणा","Small Cooler ":"छोटा कुलर","Freel":"फ्रिल","Halad Chair":"हळद खुर्ची","Sofa ":"सोफा","Focus":"फोकस","Chaurang":"चौरंग","Parda":"पडदा","Tup":"तुप","Asanpatti":"आसनपट्टी","VIP Chair ":"व्हीआयपी खुर्ची","Chair Cover":"खुर्ची कव्हर","Maharaja Sofa ":"महाराजा सोफा","Stage Table":"स्टेज टेबल","Ushi ":"उशी","Chadar":"चादर","Blanket":"ब्लँकेट","Foam Gadi":"फोम गादी","Two way ":"टू वे","Four Way":"फोर वे","Amp Pinto":"अँप पिंटो","Amp Ahuja ":"अँप अहुजा","Halad Tup":"हळद तूप","Shafing Dish":"शेफिंग डिश","Matt":"मॅट","Hawan Kund":"हवन कुंड","Dish":"डिश","Podium":"पोडियम","Stage Frame":"स्टेज फ्रेम","Buffer":"बफर","Grinder":"ग्राइंडर","Black Freel":"काळा फ्रिल","Thermos":"थर्मास","PIPE":"पाईप","Gadi Set":"गादी सेट","Table Cover":"टेबल कव्हर","Satranji":"सतरंजी","Opening Balance":"ओपनिंग","Jumbo Cooler":"जंबो कुलर"}
seen=set(); ro=3
for r in range(2,sI.max_row+1):
    iid=sI.cell(r,1).value
    if iid is None or str(iid).strip()=="": continue
    iid=str(iid).strip()
    if iid in seen: continue
    seen.add(iid)
    if ro>202: break
    nm=str(sI.cell(r,2).value or "").strip()
    wi.cell(row=ro,column=1,value=iid)
    wi.cell(row=ro,column=2,value=nm)
    wi.cell(row=ro,column=3,value=MR.get(str(sI.cell(r,2).value or ""),""))
    try: wi.cell(row=ro,column=4,value=float(sI.cell(r,4).value or 0)).number_format=MF
    except Exception: wi.cell(row=ro,column=4,value=0)
    for cx,v in [(5,sI.cell(r,3).value),(6,sI.cell(r,8).value)]:
        try: wi.cell(row=cx if False else ro,column=cx,value=float(v) if v not in (None,"") else None)
        except Exception: wi.cell(row=ro,column=cx,value=None)
    ro+=1
print("items",ro-3)
ro=3
for r in range(2,sC.max_row+1):
    shop=sC.cell(r,1).value
    if shop is None or str(shop).strip()=="": continue
    if ro>202: break
    wc.cell(row=ro,column=1,value=str(shop).strip())
    wc.cell(row=ro,column=2,value=str(sC.cell(r,2).value or "").strip())
    wc.cell(row=ro,column=3,value=str(sC.cell(r,4).value or "").strip() if sC.cell(r,4).value else "")
    wc.cell(row=ro,column=4,value=str(sC.cell(r,3).value or "").strip())
    wc.cell(row=ro,column=5,value=str(sC.cell(r,9).value or "").strip() if sC.max_column>=9 else "")
    wc.cell(row=ro,column=6,value=0).number_format=MF
    ro+=1
print("cust",ro-3)
